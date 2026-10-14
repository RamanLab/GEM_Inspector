"""Downloadable reports for a Compare Models result: an Excel workbook and a
self-contained HTML page. Both contain the naming-convention comparison, the
summary counts, and the full shared / only-in-A / only-in-B lists for genes,
metabolites, reactions and biomass components (ids with MetaNetX ids).
"""
import io
from datetime import datetime
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SECTIONS = (
    ("genes", "Genes"),
    ("metabolites", "Metabolites"),
    ("reactions", "Reactions"),
    ("biomass", "Biomass components"),
)

GROUPS = (
    ("shared", "In both models"),
    ("only_a", "Only in model A"),
    ("only_b", "Only in model B"),
)

# Subset of "shared": paired items whose flux bounds or coefficients differ.
# Listed on its own in the HTML report; in the workbook it is the "Differs" column.
DIFFER_GROUP = ("differ", "In both, but bounds or coefficients differ")

_FONT = "Arial"


def section_columns(section_key, group):
    """[(header, row -> value)] for one list. Biomass lists also show the
    coefficient in each model and whether it differs."""
    id_a = ("Model A id", lambda r: r.get("id_a", ""))
    id_b = ("Model B id", lambda r: r.get("id_b", ""))
    tail = [("Name", lambda r: r.get("name", ""))]
    if section_key != "genes":
        tail.insert(0, ("MetaNetX id", lambda r: r.get("metanetx_id", "")))
    if group == "only_a":
        return [id_a, *tail]
    if group == "only_b":
        return [id_b, *tail]
    cols = [id_a, id_b, *tail, ("Matched by", lambda r: r.get("matched_by", ""))]
    if section_key == "reactions":
        cols += [
            ("Lower bound A", lambda r: r.get("lb_a")), ("Upper bound A", lambda r: r.get("ub_a")),
            ("Lower bound B", lambda r: r.get("lb_b")), ("Upper bound B", lambda r: r.get("ub_b")),
            ("Differs", lambda r: "yes" if r.get("differences") else ""),
            ("Differences", lambda r: r.get("differences", "")),
        ]
    if section_key == "biomass":
        cols += [
            ("Coefficient A", lambda r: r.get("coefficient_a")),
            ("Coefficient B", lambda r: r.get("coefficient_b")),
            ("Differs", lambda r: "yes" if r.get("coefficient_a") != r.get("coefficient_b") else ""),
        ]
    return cols


def _consistency_text(consistent):
    return {True: "consistent", False: "differs"}.get(consistent, "can't tell")


def _naming_cell_text(info):
    text = info["label"]
    if info.get("example"):
        text += f" [example: {info['example']}; {info['coverage'] * 100:.0f}% of ids]"
    return text


def _timestamp():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


# --- Excel -------------------------------------------------------------------

def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=_FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(vertical="center")


def _autosize(ws, max_width=60):
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, 10), max_width)


def _write_summary_sheet(wb, report):
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Model comparison"
    ws["A1"].font = Font(name=_FONT, bold=True, size=14)
    ws["A2"], ws["B2"] = "Model A (loaded)", report["labels"]["a"]
    ws["A3"], ws["B3"] = "Model B", report["labels"]["b"]
    ws["A4"], ws["B4"] = "Generated", _timestamp()

    ws["A6"] = "Naming conventions"
    ws["A6"].font = Font(name=_FONT, bold=True)
    for c, h in enumerate(("", "Model A", "Model B", "Consistent?"), start=1):
        ws.cell(row=7, column=c, value=h)
    _style_header(ws, 7, 4)
    row = 8
    for r in report["naming"]["rows"]:
        for c, v in enumerate((r["aspect"], _naming_cell_text(r["model_a"]), _naming_cell_text(r["model_b"]),
                               _consistency_text(r["consistent"])), start=1):
            ws.cell(row=row, column=c, value=v)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Component counts").font = Font(name=_FONT, bold=True)
    row += 1
    for c, h in enumerate(("", "In model A", "Shared", "Only in A", "Only in B", "In model B", "Shared, but differ"), start=1):
        ws.cell(row=row, column=c, value=h)
    _style_header(ws, row, 7)
    row += 1
    for key, title in SECTIONS:
        counts = report[key].get("counts")
        if counts:
            for c, v in enumerate((title, counts["total_a"], counts["shared"], counts["only_a"],
                                   counts["only_b"], counts["total_b"], counts.get("differ", "")), start=1):
                ws.cell(row=row, column=c, value=v)
            row += 1

    bio = report["biomass"]
    row += 1
    ws.cell(row=row, column=1, value="Biomass reaction A")
    ws.cell(row=row, column=2, value=f"{bio['reaction_a']['id']} ({bio['reaction_a']['n_components']} metabolites)" if bio.get("reaction_a") else "not found")
    ws.cell(row=row + 1, column=1, value="Biomass reaction B")
    ws.cell(row=row + 1, column=2, value=f"{bio['reaction_b']['id']} ({bio['reaction_b']['n_components']} metabolites)" if bio.get("reaction_b") else "not found")

    for r in ws.iter_rows():
        for cell in r:
            if cell.font.name != _FONT:
                cell.font = Font(name=_FONT, bold=cell.font.bold, color=cell.font.color, size=cell.font.size)
    _autosize(ws, 70)


def _write_section_sheet(wb, report, key, title):
    data = report[key]
    ws = wb.create_sheet(title[:31])
    if not data.get("counts"):
        ws["A1"] = data.get("error", "Not available")
        return
    headers = ["Group"] + [h for h, _ in section_columns(key, "shared")]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for group, label in GROUPS:
        cols = section_columns(key, group)
        for item in data[group]:
            values = {h: fn(item) for h, fn in cols}
            ws.append([label] + [values.get(h, "") for h in headers[1:]])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = Font(name=_FONT)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)


def build_xlsx(report):
    """The comparison as an .xlsx: a Summary sheet plus one filterable sheet
    per component (a Group column separates shared / only in A / only in B)."""
    wb = Workbook()
    _write_summary_sheet(wb, report)
    for key, title in SECTIONS:
        _write_section_sheet(wb, report, key, title)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- HTML --------------------------------------------------------------------

_HTML_STYLE = """
body{font:14px/1.45 Arial,Helvetica,sans-serif;color:#18212f;background:#f5f7fb;margin:0}
main{max-width:1200px;margin:0 auto;padding:28px 20px 60px}
h1{margin:0 0 4px}h2{margin:34px 0 10px;border-bottom:2px solid #e5eaf1;padding-bottom:4px}
.meta{color:#64748b;margin-bottom:18px}
table{border-collapse:collapse;width:100%;background:#fff;margin-bottom:12px}
th,td{border:1px solid #e5eaf1;padding:6px 9px;text-align:left;vertical-align:top;font-size:13px}
th{background:#1f2937;color:#fff}
.good{color:#047857;font-weight:700}.bad{color:#b91c1c;font-weight:700}.warn{color:#b45309;font-weight:700}
details{background:#fff;border:1px solid #e5eaf1;border-radius:8px;margin:8px 0;padding:8px 12px}
summary{cursor:pointer;font-weight:700}
.small{color:#64748b;font-size:12px}
nav a{margin-right:14px}
"""


def _table(headers, rows):
    head = "".join(f"<th>{escape(str(h))}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{escape('' if v is None else str(v))}</td>" for v in row) + "</tr>" for row in rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _naming_html(report):
    cls = {True: "good", False: "bad", None: "warn"}
    rows = "".join(
        f"<tr><td>{escape(r['aspect'])}</td><td>{escape(_naming_cell_text(r['model_a']))}</td>"
        f"<td>{escape(_naming_cell_text(r['model_b']))}</td>"
        f"<td class=\"{cls[r['consistent']]}\">{_consistency_text(r['consistent'])}</td></tr>"
        for r in report["naming"]["rows"])
    return ("<table><thead><tr><th></th><th>Model A</th><th>Model B</th><th>Consistent?</th></tr></thead>"
            f"<tbody>{rows}</tbody></table>")


def _summary_html(report):
    rows = [(title, c["total_a"], c["shared"], c["only_a"], c["only_b"], c["total_b"], c.get("differ", ""))
            for key, title in SECTIONS if (c := report[key].get("counts"))]
    return _table(["", "In model A", "Shared", "Only in A", "Only in B", "In model B", "Shared, but differ"], rows)


def _section_html(report, key, title):
    data = report[key]
    parts = [f'<h2 id="{key}">{escape(title)}</h2>']
    if key == "biomass":
        for side in ("a", "b"):
            r = data.get(f"reaction_{side}")
            parts.append(f"<div><strong>Model {side.upper()}:</strong> "
                         + (f"{escape(r['id'])} <span class=\"small\">{escape(r['name'])} · {r['n_components']} metabolites</span>" if r else "not found")
                         + "</div>")
    if not data.get("counts"):
        parts.append(f"<p class=\"bad\">{escape(data.get('error', 'Not available'))}</p>")
        return "".join(parts)
    groups = GROUPS + ((DIFFER_GROUP,) if "differ" in data else ())
    for group, label in groups:
        cols = section_columns(key, "shared" if group == "differ" else group)
        rows = [[fn(item) for _, fn in cols] for item in data[group]]
        body = _table([h for h, _ in cols], rows) if rows else "<p class=\"small\">Nothing in this group.</p>"
        parts.append(f"<details><summary>{escape(label)} ({len(rows)})</summary>{body}</details>")
    return "".join(parts)


def build_html(report):
    """The comparison as one self-contained HTML page."""
    toc = "".join(f'<a href="#{key}">{title}</a>' for key, title in SECTIONS)
    sections = "".join(_section_html(report, key, title) for key, title in SECTIONS)
    labels = report["labels"]
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Model comparison: {escape(labels['a'])} vs {escape(labels['b'])}</title>
<style>{_HTML_STYLE}</style></head><body><main>
<h1>Model comparison</h1>
<div class="meta">Model A (loaded): <strong>{escape(labels['a'])}</strong> &middot; Model B: <strong>{escape(labels['b'])}</strong> &middot; generated {_timestamp()}</div>
<nav>{toc}</nav>
<h2>Naming conventions</h2>{_naming_html(report)}
<h2>Summary</h2>{_summary_html(report)}
{sections}
</main></body></html>"""
