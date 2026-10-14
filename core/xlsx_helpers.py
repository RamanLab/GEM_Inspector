"""Small openpyxl helpers shared by the Excel report builders."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADER_FILL = "1F2937"


def write_table(ws, headers, rows, start_row=1, widths_cap=60, wrap_columns=()):
    """Write a header row and data rows, style them, freeze the header and
    add a filter. `wrap_columns` are header names whose text wraps."""
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=FONT)
            if headers[c - 1] in wrap_columns:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
    autosize(ws, widths_cap)


def autosize(ws, max_width=60):
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, 10), max_width)


def title_cell(ws, ref, text, size=14):
    ws[ref] = text
    ws[ref].font = Font(name=FONT, bold=True, size=size)


def workbook_bytes(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def new_workbook():
    return Workbook()
